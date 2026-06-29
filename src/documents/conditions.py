"""Génération des Conditions Particulières de Vente T1 / T2 (Les Menus Services).

CLAUDE.md §4 :
- T1 = tournée Lundi/Jeudi  → modèle « Conditions ... T1 ... »
- T2 = tournée Mardi/Vendredi → modèle « Conditions ... T2 ... - Copie »
Le paramètre `tournee` choisit le bon modèle de base. On NE reconstruit PAS le
tableau du planning de livraison (Table 1, partie réglementaire qui diffère
T1/T2) : on part du bon modèle et on remplit uniquement les trous.

Champs remplis (section finale + identité) :
- identité client/bénéficiaire, lieu de prestation (si différent) ;
- commencement d'exécution (attendre 14 j OU démarrer avant) + 1re livraison ;
- jours de repas : cellules-marques de la ligne des jours (Table 0) ;
- coût des prestations : MÊMES valeurs que le devis, RÉINJECTÉES (pas de
  recalcul). La trame CPV n'a PAS de ligne « mensuel » ni « après crédit
  d'impôt » : ces données restent sur le devis (décision client). On ne remplit
  que ce qui existe dans la trame officielle.
- « Fait à … le … ».

Méthode : copie-template balisée (docxtpl), originaux jamais modifiés.
Repérage des paragraphes par TEXTE (pas par index) car T1/T2 sont décalés.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docxtpl import DocxTemplate

from src.calcul import Resultat
from src.documents.devis import (  # mêmes rendus « formule » / tarifs que le devis
    _detail_formules,
    tarifs_cumules,
)

from src.chemins import RACINE
_MODELES = {
    "T1": RACINE / "modeles" / "Conditions Particulieres de Vente T1 à partir du 15 juin Menus Services Orange.docx",
    "T2": RACINE / "modeles" / "Conditions Particulieres de Vente T2 à partir du 15 juin Menus Services Orange - Copie.docx",
}
_TEMPLATES = {
    "T1": RACINE / "modeles" / "_template_conditions_T1.docx",
    "T2": RACINE / "modeles" / "_template_conditions_T2.docx",
}

JOURS = ["Lundi", "Mardi", "Mercredi", "Jeudi", "Vendredi", "Samedi", "Dimanche"]


def _cle_jour(nom: str) -> str:
    """'Lundi' -> 'lundi' (clé de balise jour_<x>)."""
    return (
        nom.strip().lower()
        .replace("é", "e").replace("è", "e").replace("â", "a")
    )


@dataclass
class InfosConditions:
    """Champs non tarifaires des CPV (alimentés par l'UI plus tard)."""

    client_nom: str = ""
    client_adresse: str = ""
    benef_nom: str = ""
    benef_prenom: str = ""
    benef_adresse: str = ""
    tuteur: str = ""                   # « Prénom NOM (Organisme) » si tutelle
    lieu_prestation: str = ""          # adresse de réalisation si différente
    commencement: str = "attendre"     # "attendre" (14 j) ou "avant"
    date_premiere_livraison: str = ""
    jours_repas: list[str] = field(default_factory=list)  # sous-ensemble de JOURS
    mode_paiement: str = ""            # "prelevement", "virement", "cheque" ou ""
    lieu: str = ""                     # « Fait à … »
    date: str = ""


# --------------------------------------------------------------------------- #
# Helpers d'édition de paragraphe
# --------------------------------------------------------------------------- #
def _ecrire_paragraphe(paragraphe, texte: str) -> None:
    """Écrit le texte dans le 1er run (mise en forme du paragraphe conservée)."""
    runs = paragraphe.runs
    if not runs:
        paragraphe.add_run(texte)
        return
    runs[0].text = texte
    for r in runs[1:]:
        r.text = ""


def _ecrire_en_preservant_indentation(paragraphe, coeur: str) -> None:
    """Réécrit un paragraphe en conservant ses espaces de tête (indentation)."""
    txt = paragraphe.text
    tete = txt[: len(txt) - len(txt.lstrip(" "))]
    _ecrire_paragraphe(paragraphe, tete + coeur)


# Règles : (prédicat sur le texte STRIPPÉ, cœur de remplacement avec balises).
# Le cœur reprend le libellé réglementaire ; seuls les pointillés deviennent
# des balises. L'indentation de tête est préservée séparément.
_REGLES = [
    (lambda t: t.startswith("NOM, Prénom"),
     "NOM, Prénom : {{ client_nom }}"),
    (lambda t: t.startswith("Adresse du client"),
     "Adresse du client : {{ client_adresse }}"),
    (lambda t: t.startswith("Nom ") and "Prénom" in t,
     "Nom {{ benef_nom }}  Prénom {{ benef_prenom }}"),
    (lambda t: t.startswith("Adresse")
        and not t.startswith("Adresse du client")
        and not t.startswith("Adresse de la réalisation"),
     "Adresse {{ benef_adresse }}"),
    (lambda t: t.startswith("Jour prévu pour la première livraison"),
     "Jour prévu pour la première livraison : {{ date_premiere_livraison }}"),
    (lambda t: t.startswith("Formule choisie"),
     "Formule choisie : {{ formule_detail }}"),
    (lambda t: t.startswith("Tarif global"),
     "Tarif global de la formule de repas choisie : {{ tarif_ttc }} euros TTC ({{ tarif_ht }} euros HT)"),
    (lambda t: t.startswith("- dont Prix du repas"),
     "- dont Prix du repas {{ repas_ttc }} euros TTC (TVA 10% - {{ repas_ht }} euros HT)"),
    (lambda t: t.startswith("- dont prix du Service"),
     "- dont prix du Service livraison d’un repas {{ service_ttc }} euros TTC (TVA 10% - {{ service_ht }} euros HT)"),
    (lambda t: t.startswith("Fait à"),
     "Fait à {{ lieu }} le {{ date }}, en deux exemplaires."),
]


#: Marqueur mc:AlternateContent (formes dessinées avec repli image/VML).
_MC_ALTERNATE = "{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent"


def _supprimer_cases_dessinees(paragraphe) -> None:
    """Supprime les petites cases DESSINÉES (≈3 mm) ancrées sur ce paragraphe.

    Comme la marque « X » est écrite en TEXTE devant l'option choisie (même
    style que la ligne « X Je souhaite… » de la trame), ces cases flottantes
    ne serviraient qu'à créer un X « à côté » d'une case vide — défaut signalé
    par le client. Les grands cadres de signature (ancrés ailleurs) ne sont
    pas touchés.
    """
    for el in list(paragraphe._p.iter(_MC_ALTERNATE)):
        el.getparent().remove(el)


def _supprimer_puce(paragraphe) -> None:
    """Retire la puce de liste (numPr) d'un paragraphe, sans toucher au texte.

    Dans les trames, certaines options à cocher portent une puce « ⬜ »
    (U+2B1C, liste numId 1) qui joue le rôle de case à cocher. La marque « X »
    étant écrite en texte, la puce-case resterait VIDE à côté du X — on la
    retire des seules lignes d'option réécrites. L'indentation du paragraphe
    (w:ind) est conservée.
    """
    pPr = paragraphe._p.pPr
    if pPr is not None:
        numPr = pPr.find(qn("w:numPr"))
        if numPr is not None:
            pPr.remove(numPr)


def _baliser_paiement(doc) -> None:
    """Balise les deux lignes « Modalités de paiement » (marque devant l'option).

    Les « cases » de la trame sont hétérogènes (forme flottante devant
    « Prélèvement » en T1, puce ⬜ en T2 ; puce devant « Virement » ; forme
    flottante devant « Chèque ») : on les retire toutes et on marque l'option
    choisie d'un X texte, comme la section « Commencement d'exécution ».
    """
    for p in doc.paragraphs:
        t = p.text.strip()
        if t.startswith("Prélèvement automatique"):
            _supprimer_cases_dessinees(p)
            _supprimer_puce(p)
            _ecrire_paragraphe(
                p,
                "{{ marque_prelevement }}   Prélèvement automatique "
                "(fournir RIB et autorisation de prélèvement automatique)",
            )
        elif t.startswith("Virement"):
            _supprimer_cases_dessinees(p)
            _supprimer_puce(p)
            _ecrire_paragraphe(
                p,
                "{{ marque_virement }}   Virement             "
                "{{ marque_cheque }}   Chèque bancaire",
            )


def _baliser_tutelle(doc) -> None:
    """Balise la ligne « représenté par ……, tuteur » du bloc tutelle officiel.

    La trame prévoit déjà la clause « Si le client a été placé sous tutelle… »
    suivie de « représenté par ……, tuteur » : on remplit ces pointillés (le
    rendu garde « …… » si le client n'est pas sous tutelle, comme l'original).
    """
    for p in doc.paragraphs:
        t = p.text.strip()
        if t.startswith("représenté par") and t.endswith("tuteur"):
            _ecrire_en_preservant_indentation(p, "représenté par {{ tuteur_detail }}, tuteur")
            return


def _baliser_lieu(doc) -> None:
    """Balise la ligne de valeur sous « Adresse de la réalisation… » (Article 2).

    Le bloc reste TOUJOURS présent : le rendre conditionnel supprimait
    l'Article 2 entier et décalait la numérotation automatique de tous les
    articles suivants (retour client). Le rendu garde des pointillés si aucun
    lieu distinct n'est saisi, comme la trame vierge.
    """
    paras = doc.paragraphs
    for i, p in enumerate(paras):
        if p.text.strip().startswith("Adresse de la réalisation"):
            valeur = paras[i + 1] if i + 1 < len(paras) else None
            if valeur is not None and "…" in valeur.text:
                _ecrire_paragraphe(valeur, "{{ lieu_prestation_aff }}")
            else:
                p.runs[0].text = p.runs[0].text.rstrip() + " : {{ lieu_prestation_aff }}"
            return


# --------------------------------------------------------------------------- #
# Construction de la copie-template balisée
# --------------------------------------------------------------------------- #
def construire_template(tournee: str) -> Path:
    """Crée (ou recrée) la copie-template balisée pour la tournée T1/T2."""
    if tournee not in _MODELES:
        raise ValueError(f"Tournée inconnue : {tournee!r} (attendu 'T1' ou 'T2')")
    doc = Document(str(_MODELES[tournee]))

    for p in doc.paragraphs:
        t = p.text.strip()
        if not t:
            continue

        # Le bordereau de rétractation détachable (en fin de document) est un
        # formulaire que le client remplit À LA MAIN s'il se rétracte : il doit
        # rester VIERGE. On arrête tout remplissage dès qu'on l'atteint.
        if t.startswith("Bordereau susceptible") or t.startswith("Formulaire de rétractation"):
            break

        # Lignes à remplacement standard (identité, coût, fait à…).
        for predicat, coeur in _REGLES:
            if predicat(t):
                _ecrire_en_preservant_indentation(p, coeur)
                break
        else:
            # Commencement — option « attendre 14 j » : la marque « X » figée
            # devient conditionnelle, sans toucher au texte réglementaire.
            if "expiration du délai" in t and "14 jours" in t:
                p.runs[0].text = re.sub(r"X", "{{ marque_attendre }}", p.runs[0].text, count=1)
            # Commencement — option « démarrer avant » : on préfixe la marque
            # et on retire la puce-case « ⬜ » (le X texte la remplace, sinon
            # la croix apparaît À CÔTÉ d'une case vide — défaut signalé).
            elif "démarre avant la fin" in t:
                p.runs[0].text = "{{ marque_avant }}   " + p.runs[0].text
                _supprimer_puce(p)

    # --- Lieu de prestation (Article 2, toujours présent) --------------------
    # Le libellé réglementaire (avec sa coquille « l'adresser ») n'est PAS
    # modifié ; seule la ligne de pointillés sous le titre est balisée.
    _baliser_lieu(doc)
    _baliser_paiement(doc)
    _baliser_tutelle(doc)

    # --- Pagination : « Fait à … » + signatures + bordereau sur leur page ----
    # Dans la trame, ce bloc occupe la dernière page « par chance » (le texte
    # remplit pile les pages 1-2). Les lignes conditionnelles supprimées au
    # rendu décalent tout et coupaient le bloc signatures en deux. Un saut de
    # page explicite avant « Fait à » fige la pagination de l'original. Les
    # paragraphes vides juste avant sont retirés : sinon ils peuvent rester
    # seuls sur une page (page blanche) quand le texte remplit pile la page.
    paras = doc.paragraphs
    for i, p in enumerate(paras):
        if p.text.strip().startswith("Fait à"):
            p.paragraph_format.page_break_before = True
            j = i - 1
            while j >= 0 and not paras[j].text.strip():
                paras[j]._p.getparent().remove(paras[j]._p)
                j -= 1
            break

    # --- Ligne des jours (Table 0) : remplir les cellules-marques ------------
    # La marque est CENTRÉE dans sa cellule bordée (la « case ») ; la trame
    # d'origine la décalait avec des espaces, ce qui collait le X au bord.
    for table in doc.tables:
        cells = table.rows[0].cells
        if cells and cells[0].text.strip() == "Lundi":
            for i in range(0, len(cells) - 1, 2):
                jour = cells[i].text.strip()
                if jour in JOURS:
                    para = cells[i + 1].paragraphs[0]
                    _ecrire_paragraphe(para, f"{{{{ jour_{_cle_jour(jour)} }}}}")
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            break

    dst = _TEMPLATES[tournee]
    doc.save(str(dst))
    return dst


# --------------------------------------------------------------------------- #
# Contexte de rendu
# --------------------------------------------------------------------------- #
def contexte_conditions(infos: InfosConditions, resultat: Resultat) -> dict:
    """Construit le contexte docxtpl. Valeurs tarifaires = celles du devis."""
    selection = {_cle_jour(j) for j in infos.jours_repas}
    ctx = {
        # identité (Article 1 TOUJOURS rempli — identité client si identique)
        "client_nom": infos.client_nom,
        "client_adresse": infos.client_adresse,
        "benef_nom": infos.benef_nom,
        "benef_prenom": infos.benef_prenom,
        "benef_adresse": infos.benef_adresse,
        # bloc tutelle officiel : pointillés conservés si pas de tuteur
        "tuteur_detail": infos.tuteur.strip() or "……",
        # Article 2 : pointillés conservés si pas de lieu distinct
        "lieu_prestation_aff": infos.lieu_prestation.strip()
        or "……………………………………………………………………………………",
        # commencement (marque « X » sur l'option choisie)
        "marque_attendre": "X" if infos.commencement == "attendre" else "",
        "marque_avant": "X" if infos.commencement == "avant" else "",
        "date_premiere_livraison": infos.date_premiere_livraison,
        # modalités de paiement (marque « X » sur le mode choisi)
        "marque_prelevement": "X" if infos.mode_paiement == "prelevement" else "",
        "marque_virement": "X" if infos.mode_paiement == "virement" else "",
        "marque_cheque": "X" if infos.mode_paiement == "cheque" else "",
        # coût : MÊMES tarifs AU REPAS cumulés que le devis
        "formule_detail": _detail_formules(resultat),
        **tarifs_cumules(resultat),
        # « Fait à … le … »
        "lieu": infos.lieu,
        "date": infos.date,
    }
    # jours de repas : « X » si coché, vide sinon
    for j in JOURS:
        cle = _cle_jour(j)
        ctx[f"jour_{cle}"] = "X" if cle in selection else ""
    return ctx


def generer_conditions(
    tournee: str,
    infos: InfosConditions,
    resultat: Resultat,
    sortie: Path,
    template: Optional[Path] = None,
) -> Path:
    """Génère les CPV (T1 ou T2) remplies dans `sortie`."""
    if template is None:
        template = _TEMPLATES[tournee]
    if not Path(template).exists():
        construire_template(tournee)

    tpl = DocxTemplate(str(template))
    tpl.render(contexte_conditions(infos, resultat))
    sortie.parent.mkdir(parents=True, exist_ok=True)
    tpl.save(str(sortie))
    return sortie
