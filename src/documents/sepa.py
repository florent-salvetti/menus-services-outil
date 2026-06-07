"""Génération du mandat de prélèvement SEPA (Les Menus Services).

⚠️ RGPD (CLAUDE.md §9) : IBAN/BIC NON PERSISTÉS. Ils transitent uniquement en
mémoire le temps de générer le document, puis sont écartés par l'appelant.
Aucune écriture en base ici.

Particularité de la trame (Autorisation de prélèvements.docx) :
- C'est une autorisation ANCIEN FORMAT : la section « compte à débiter » utilise
  la grille RIB française classique (Établissement / Guichet / N° de compte /
  Clé), et NE comporte PAS de champ IBAN ni BIC.
- Choix validé client (Option 1) : on DÉCOMPOSE l'IBAN français dans cette
  grille (déterministe), après double validation (clé IBAN mod 97 + clé RIB).
  Le BIC n'est donc pas inscrit sur le document (pas de champ pour lui).
  → À confirmer avec le client : ce format de mandat est-il toujours accepté
    par sa banque, ou faut-il migrer vers un mandat SEPA moderne (IBAN/BIC) ?
- Les champs apparaissent en double (2 exemplaires) : on remplit les deux.
- ICS (FR18ZZZ86BB45) et créancier (SAS BON A DOM') sont déjà dans la trame :
  on n'y touche pas.

Méthode : copie-template balisée (docxtpl), original jamais modifié.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Optional

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.text.paragraph import Paragraph
from docxtpl import DocxTemplate

from src.iban import RibFr, valider_pour_mandat

RACINE = Path(__file__).resolve().parent.parent.parent
MODELE_ORIGINAL = RACINE / "modeles" / "Autorisation de prélèvements.docx"
TEMPLATE_GENERE = RACINE / "modeles" / "_template_sepa.docx"


@dataclass
class InfosMandat:
    """Champs non bancaires du mandat (alimentés par l'UI plus tard)."""

    debiteur_nom: str = ""       # = client en général
    debiteur_adresse: str = ""
    banque_nom: str = ""
    banque_adresse: str = ""
    lieu_signature: str = ""
    date_signature: str = ""


# --------------------------------------------------------------------------- #
# Construction de la copie-template balisée
# --------------------------------------------------------------------------- #
def _ecrire_paragraphe(paragraphe: Paragraph, texte: str) -> None:
    """Écrit le texte dans un paragraphe en conservant sa mise en forme.

    Réutilise la police du 1er run s'il existe ; sinon ajoute un run.
    """
    runs = paragraphe.runs
    if not runs:
        paragraphe.add_run(texte)
        return
    runs[0].text = texte
    for r in runs[1:]:
        r.text = ""


def _paragraphes_zone(box) -> list[Paragraph]:
    return [Paragraph(p, None) for p in box.iter(qn("w:p"))]


def _cellule_xml(texte: str, largeur: int, *, taille: int, gras: bool) -> str:
    """Génère le XML d'une cellule de tableau centrée (police `taille` 1/2-pt)."""
    rpr = f'<w:sz w:val="{taille}"/>' + ("<w:b/>" if gras else "")
    return (
        f'<w:tc><w:tcPr><w:tcW w:w="{largeur}" w:type="dxa"/>'
        f'<w:vAlign w:val="center"/></w:tcPr>'
        f'<w:p><w:pPr><w:jc w:val="center"/><w:spacing w:after="0"/>'
        f"<w:rPr>{rpr}</w:rPr></w:pPr>"
        f"<w:r><w:rPr>{rpr}</w:rPr><w:t>{texte}</w:t></w:r></w:p></w:tc>"
    )


# Colonnes (twips) — total 3300, tient dans la boîte de 3600 twips (180 pt).
_COLS = (760, 700, 1440, 400)
_ENTETES = ("Établiss.", "Guichet", "N° de compte", "Clé")
_VALEURS = ("{{ rib_banque }}", "{{ rib_guichet }}", "{{ rib_compte }}", "{{ rib_cle }}")


def _table_compte():
    """Construit un tableau bordé 2×4 (en-têtes + valeurs balisées) pour le RIB.

    Chaque valeur dans sa cellule => alignement garanti sous son intitulé,
    indépendamment de la police (proportionnelle). Reconstruit uniquement la
    zone « compte à débiter », sans toucher au reste de la trame.
    """
    bord = '<w:{0} w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
    bordures = "".join(
        bord.format(b) for b in ("top", "left", "bottom", "right", "insideH", "insideV")
    )
    grid = "".join(f'<w:gridCol w:w="{w}"/>' for w in _COLS)
    entetes = "".join(
        _cellule_xml(t, w, taille=13, gras=False) for t, w in zip(_ENTETES, _COLS)
    )
    valeurs = "".join(
        _cellule_xml(v, w, taille=18, gras=True) for v, w in zip(_VALEURS, _COLS)
    )
    xml = (
        f"<w:tbl {nsdecls('w')}>"
        f'<w:tblPr><w:tblW w:w="3300" w:type="dxa"/><w:jc w:val="center"/>'
        f"<w:tblBorders>{bordures}</w:tblBorders>"
        f'<w:tblLayout w:type="fixed"/></w:tblPr>'
        f"<w:tblGrid>{grid}</w:tblGrid>"
        f"<w:tr>{entetes}</w:tr><w:tr>{valeurs}</w:tr>"
        f"</w:tbl>"
    )
    return parse_xml(xml)


#: Balise mc:AlternateContent (formes Office avec repli VML).
_MC_ALTERNATE = "{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent"


def _supprimer_grille_residuelle(doc) -> int:
    """Supprime les traits dessinés de l'ANCIENNE grille RIB à cases.

    Ces traits sont des formes « ligne » (DrawingML `straightConnector1`, repli
    VML `v:line`) encapsulées dans `mc:AlternateContent`. Elles ne servent plus
    depuis qu'on remplit le compte via un vrai tableau. On ne retire QUE les
    formes ligne (jamais les zones de texte = `txbxContent`). Retourne le nombre
    d'éléments supprimés.
    """
    from lxml import etree

    a_supprimer = []
    for el in doc.element.body.iter(_MC_ALTERNATE):
        s = etree.tostring(el).decode("utf-8", "ignore")
        if ("straightConnector1" in s or "v:line" in s) and "txbxContent" not in s:
            a_supprimer.append(el)
    for el in a_supprimer:
        el.getparent().remove(el)
    return len(a_supprimer)


def construire_template(
    src: Path = MODELE_ORIGINAL, dst: Path = TEMPLATE_GENERE
) -> Path:
    """Crée (ou recrée) la copie-template balisée du mandat SEPA."""
    doc = Document(str(src))

    # Retire les traits résiduels de l'ancienne grille RIB (remplacée par un
    # tableau). À faire avant la sauvegarde, sans toucher aux zones de texte.
    _supprimer_grille_residuelle(doc)

    # --- Zones de texte (txbxContent), en double (2 exemplaires) -------------
    for box in doc.element.body.iter(qn("w:txbxContent")):
        paras = _paragraphes_zone(box)
        textes = [p.text for p in paras]
        entete = " ".join(textes)

        if "DEBITEUR" in entete:
            # 1re ligne vide après le label.
            cible = next(p for p in paras[1:] if not p.text.strip())
            _ecrire_paragraphe(cible, "{{ debiteur }}")

        elif "ETABLISSEMENT FINANCIER" in entete:
            cible = next(p for p in paras[1:] if not p.text.strip())
            _ecrire_paragraphe(cible, "{{ etablissement }}")

        elif "COMPTE A DEBITER" in entete:
            # On garde le titre « COMPTE A DEBITER », on remplace les en-têtes
            # espacés + la ligne vide par un vrai tableau aligné.
            enfants = box.findall(qn("w:p"))
            for p in enfants[1:]:  # tout sauf le titre
                box.remove(p)
            box.append(_table_compte())

    # --- Corps : ligne « Date : » -------------------------------------------
    for p in doc.paragraphs:
        if p.text.strip().startswith("Date"):
            _ecrire_paragraphe(p, "Date : {{ date_signature }} à {{ lieu_signature }}")

    doc.save(str(dst))
    return dst


# --------------------------------------------------------------------------- #
# Contexte de rendu
# --------------------------------------------------------------------------- #
def contexte_mandat(infos: InfosMandat, rib: RibFr) -> dict:
    """Construit le dict de rendu docxtpl (\\n -> saut de ligne via docxtpl)."""
    debiteur = "\n".join(x for x in (infos.debiteur_nom, infos.debiteur_adresse) if x.strip())
    etablissement = "\n".join(x for x in (infos.banque_nom, infos.banque_adresse) if x.strip())
    return {
        "debiteur": debiteur,
        "etablissement": etablissement,
        # Blocs RIB, un par cellule du tableau « compte à débiter ».
        "rib_banque": rib.banque,
        "rib_guichet": rib.guichet,
        "rib_compte": rib.compte,
        "rib_cle": rib.cle_rib,
        "date_signature": infos.date_signature,
        "lieu_signature": infos.lieu_signature,
    }


# --------------------------------------------------------------------------- #
# Génération (avec validation IBAN obligatoire)
# --------------------------------------------------------------------------- #
def generer_mandat(
    infos: InfosMandat,
    iban: str,
    bic: str,
    sortie: Path,
    template: Optional[Path] = None,
) -> Path:
    """Génère le mandat SEPA rempli dans `sortie`.

    Valide l'IBAN (mod 97 + clé RIB) AVANT toute génération : si invalide,
    lève IbanInvalide et n'écrit aucun fichier.

    ⚠️ iban/bic ne sont PAS persistés : utilisés uniquement ici, en mémoire.
    (Le BIC n'a pas de champ dans cette trame ; il n'est pas inscrit.)
    """
    # Validation bloquante (lève IbanInvalide avec message précis si KO).
    rib = valider_pour_mandat(iban)

    if template is None:
        template = TEMPLATE_GENERE
    if not Path(template).exists():
        construire_template(dst=Path(template))

    tpl = DocxTemplate(str(template))
    tpl.render(contexte_mandat(infos, rib))
    sortie.parent.mkdir(parents=True, exist_ok=True)
    tpl.save(str(sortie))
    return sortie
